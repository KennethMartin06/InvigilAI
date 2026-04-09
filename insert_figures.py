#!/usr/bin/env python3
"""
insert_figures.py  v4
Replaces every [IMAGE PLACEHOLDER] in the IEEE paper docx with the correct
figure image. Uses document-order paragraph scanning (not sibling lookup) so
it works regardless of nesting (body paras, table cells, text boxes).

Run in WSL:
    cp "/mnt/c/Users/kenne/Downloads/Multi_Modal_Cheating_Detection_IEEE_Paper.docx" ~/invigilai/
    python3 insert_figures.py
"""

import re, sys, shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# ── Paths ─────────────────────────────────────────────────────────────────────
CANDIDATES = [
    Path.home() / "invigilai" / "Multi_Modal_Cheating_Detection_IEEE_Paper_FINAL.docx",
    Path.home() / "invigilai" / "Multi_Modal_Cheating_Detection_IEEE_Paper.docx",
    Path("/mnt/c/Users/kenne/Downloads/Multi_Modal_Cheating_Detection_IEEE_Paper_FINAL.docx"),
    Path("/mnt/c/Users/kenne/Downloads/Multi_Modal_Cheating_Detection_IEEE_Paper.docx"),
]
SRC    = next((c for c in CANDIDATES if c.exists()), None)
DEST   = None if SRC is None else SRC.parent / (SRC.stem + "_WITH_FIGURES.docx")
CHARTS = Path(__file__).parent / "paper_charts"

# ── Figure map ────────────────────────────────────────────────────────────────
# (regex pattern matched against the CAPTION paragraph after the placeholder,
#  image filename,  display width in inches)
#  IEEE single-col ≈ 3.3"   |   double-col ≈ 6.5"
FIGURE_MAP = [
    (r"fig\.?\s*1\b|efficientnet|mbconv",
        "paper_fig1_efficientnet_architecture.png",  6.3),
    (r"fig\.?\s*2\b|training.*accur|val.*accur",
        "paper_fig2_training_accuracy.png",          3.2),
    (r"fig\.?\s*3\b|training.*loss|val.*loss",
        "paper_fig3_training_loss.png",              3.2),
    (r"fig\.?\s*4\b|confusion.matrix",
        "paper_fig4_confusion_matrix.png",           3.0),
    (r"fig\.?\s*5\b|precision.recall",
        "paper_fig5_pr_curve.png",                   3.0),
    (r"fig\.?\s*6\b|roc curve|auc.*0\.98",
        "paper_fig6_roc_curve.png",                  3.0),
    (r"fig\.?\s*7\b|per.modality|modality.*f1|ablation",
        "paper_fig7_per_modality_f1.png",            3.2),
    (r"fig\.?\s*8\b|normal case 1|shap summary.*right",
        "paper_fig8_normal_case1.png",               3.2),
    (r"fig\.?\s*9\b|normal case 2|diffuse activation",
        "paper_fig9_normal_case2.png",               3.2),
    (r"fig\.?\s*10\b|normal case 3|eye and mouth",
        "paper_fig10_normal_case3.png",              3.2),
    (r"fig\.?\s*11\b|normal case 4|keystroke dynamics",
        "paper_fig11_normal_case4.png",              3.2),
    (r"fig\.?\s*12\b|suspicious case 1|gaze deviation|lateral gaze",
        "paper_fig12_suspicious_gaze.png",           3.2),
    (r"fig\.?\s*13\b|suspicious case 2|audio anomaly|mfcc",
        "paper_fig13_suspicious_audio.png",          3.2),
    (r"fig\.?\s*14\b|suspicious case 3|keystroke anomaly|copy.paste",
        "paper_fig14_suspicious_keystroke.png",      3.2),
    (r"fig\.?\s*15\b|monitoring dashboard|active sessions",
        "paper_fig15_dashboard.png",                 6.3),
    (r"fig\.?\s*16\b|risk score heatmap|temporal heatmap",
        "paper_fig16_risk_heatmap.png",              6.3),
    (r"fig\.?\s*17\b|alert panel|flagged sessions",
        "paper_fig17_alert_panel.png",               6.3),
    (r"fig\.?\s*18\b|session replay|evidence timeline",
        "paper_fig18_session_replay.png",            6.3),
]


def _para_text(p_elem):
    return "".join(t.text or "" for t in p_elem.iter(qn("w:t")))


def _match(text):
    t = text.lower()
    for pattern, fname, width in FIGURE_MAP:
        if re.search(pattern, t):
            p = CHARTS / fname
            if p.exists():
                return p, width
    return None


def _clear_para(p_elem):
    """Remove all run/text children from a paragraph element."""
    for child in list(p_elem):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in ("r", "hyperlink", "ins", "del", "proofErr",
                   "bookmarkStart", "bookmarkEnd"):
            p_elem.remove(child)


def _insert_image_into_para(doc, p_elem, img_path, width_in):
    """
    Clear p_elem and fill it with a centred inline image.
    Uses a temp paragraph added to body to get the correct relationship ID,
    then moves its XML content into p_elem.
    """
    # 1. Add temp paragraph at end of body
    tmp = doc.add_paragraph()
    tmp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tmp.add_run()
    run.add_picture(str(img_path), width=Inches(width_in))
    tmp_elem = tmp._element

    # 2. Steal the run element (contains the drawing/image)
    img_run = None
    for child in list(tmp_elem):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "r":
            img_run = child
            break

    if img_run is None:
        doc.element.body.remove(tmp_elem)
        return False

    # 3. Also steal the paragraph properties (centering) from tmp
    pPr = tmp_elem.find(qn("w:pPr"))

    # 4. Clear placeholder paragraph and populate with image content
    _clear_para(p_elem)

    # Set centering on placeholder para
    existing_pPr = p_elem.find(qn("w:pPr"))
    if pPr is not None:
        import copy
        new_pPr = copy.deepcopy(pPr)
        if existing_pPr is not None:
            p_elem.remove(existing_pPr)
        p_elem.insert(0, new_pPr)

    import copy
    p_elem.append(copy.deepcopy(img_run))

    # 5. Remove temp paragraph
    doc.element.body.remove(tmp_elem)
    return True


def main():
    if SRC is None:
        print("ERROR: docx not found.")
        print("  cp /mnt/c/Users/kenne/Downloads/Multi_Modal_...docx ~/invigilai/")
        sys.exit(1)

    print(f"Source : {SRC.name}")
    print(f"Output : {DEST.name}")
    print(f"Charts : {CHARTS}/\n")

    doc = Document(str(SRC))

    # ── Collect ALL paragraphs in document order ──────────────────────────────
    all_p    = list(doc.element.body.iter(qn("w:p")))
    all_text = [_para_text(p) for p in all_p]
    n        = len(all_p)

    print(f"Total paragraphs scanned: {n}")

    # ── Find placeholders and match figures ───────────────────────────────────
    jobs = []   # (placeholder_para_idx, img_path, width, [cleanup_indices])
    for i, txt in enumerate(all_text):
        if "[image placeholder]" not in txt.lower():
            continue

        # Build context from NEXT 4 paragraphs in document order (not siblings)
        context = " ".join(all_text[i+1 : i+5])

        result = _match(context)
        if result is None:
            # Try broader context (sometimes caption is further away)
            context2 = " ".join(all_text[i : i+8])
            result = _match(context2)

        if result is None:
            print(f"  [!] No match — context: {context[:80]!r}")
            continue

        img_path, width = result

        # Find cleanup indices: "Replace with actual image" and duplicate captions
        cleanup = []
        caption_seen = False
        for j in range(i+1, min(i+6, n)):
            t = all_text[j].strip().lower()
            if t.startswith("replace with"):
                cleanup.append(j)
            elif re.match(r"fig\.?\s*\d", t):
                if caption_seen:
                    cleanup.append(j)   # duplicate caption
                else:
                    caption_seen = True

        jobs.append((i, img_path, width, cleanup))

    print(f"Matched {len(jobs)} placeholder(s).\n")

    # ── Apply replacements ────────────────────────────────────────────────────
    replaced = 0
    for i, img_path, width, cleanup_idxs in jobs:
        p_elem = all_p[i]
        ok = _insert_image_into_para(doc, p_elem, img_path, width)
        if ok:
            print(f"  [OK] Fig matched → {img_path.name}  ({width}\")")
            replaced += 1
        else:
            print(f"  [FAIL] Could not insert {img_path.name}")
            continue

        # Clean up "Replace with actual image" and duplicate caption paragraphs
        for ci in cleanup_idxs:
            _clear_para(all_p[ci])

    print(f"\nReplaced: {replaced} / {len(jobs)}")

    if replaced == 0:
        print("\nWARNING: No replacements made.")
        print("Run with --debug to inspect paragraph text.")
        sys.exit(1)

    # ── Save ──────────────────────────────────────────────────────────────────
    doc.save(str(DEST))
    print(f"\nSaved : {DEST}")

    win = Path("/mnt/c/Users/kenne/Downloads") / DEST.name
    try:
        shutil.copy(str(DEST), str(win))
        print(f"Copied: {win}")
    except Exception as e:
        print(f"(Windows copy failed: {e})")

    print("\nDone.")


def debug():
    if SRC is None:
        print("docx not found"); sys.exit(1)
    doc = Document(str(SRC))
    all_p = list(doc.element.body.iter(qn("w:p")))
    print(f"Scanning {SRC.name} — {len(all_p)} paragraphs\n")
    for i, p in enumerate(all_p):
        txt = _para_text(p).strip()
        if txt and any(k in txt.lower() for k in
                       ("placeholder", "fig.", "replace with", "image")):
            print(f"  [{i:04d}] {txt[:120]!r}")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--debug":
        debug()
    else:
        main()
